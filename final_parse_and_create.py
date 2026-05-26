from docx import Document
import re
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime

docx_path = 'data/DCDEA Practice Exam 2.docx'
doc = Document(docx_path)

# Parse questions from paragraphs
print("Extracting questions from DOCX...")
questions = []
current_question = None
current_lines = []

for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        continue
    
    question_match = re.match(r'^Question\s+(\d+):', text)
    if question_match:
        if current_question is not None:
            current_question['lines'] = current_lines
            questions.append(current_question)
        
        current_question = {'num': int(question_match.group(1))}
        current_lines = [text]
        continue
    
    if current_question is not None:
        if text.startswith('Correct Anwser:') or text.startswith('Correct Answer:'):
            current_lines.append(text)
            current_question['lines'] = current_lines
            questions.append(current_question)
            current_question = None
            current_lines = []
            continue
        
        current_lines.append(text)

if current_question is not None:
    current_question['lines'] = current_lines
    questions.append(current_question)

print(f"Total questions extracted: {len(questions)}")

# Smart parser
def parse_question_final(q):
    lines = q['lines'][1:]  # Skip "Question N:" line
    
    # Find separator line index
    separator_idx = -1
    for i, line in enumerate(lines):
        if line.startswith('-'):
            separator_idx = i
            break
    
    if separator_idx == -1:
        return None
    
    # Lines before separator: question context, prompt, and options
    content_lines = lines[:separator_idx]
    
    # Lines after separator: explanation and correct answer
    explanation_lines = lines[separator_idx+1:]
    
    # Find the question prompt (usually contains "Which" or similar)
    question_idx = -1
    for i, line in enumerate(content_lines):
        if 'Which' in line or 'What' in line or 'How' in line:
            question_idx = i
            break
    
    if question_idx == -1:
        question_idx = len(content_lines) - 1
    
    # Question text = everything from start up to and including the prompt line
    question_text = ' '.join(content_lines[:question_idx+1]).strip()
    
    # Options = lines after the prompt line, filtered for real options
    option_lines = content_lines[question_idx+1:]
    options = {}
    for i, line in enumerate(option_lines):
        if line and not line.startswith('-'):
            letter = chr(65 + i)  # A, B, C, D, E, F
            options[letter] = line
    
    # Extract explanation
    explanation = ""
    for line in explanation_lines:
        if 'Correct Anwser:' in line or 'Correct Answer:' in line:
            break
        if line and not line.startswith('-') and 'Documentation' not in line and 'Reference' not in line:
            explanation += line + " "
    
    # Extract correct answer
    correct_answer = ""
    for line in explanation_lines:
        if 'Correct Anwser:' in line or 'Correct Answer:' in line:
            # Extract numbers from parentheses
            match = re.search(r'\(([^)]+)\)', line)
            if match:
                correct_answer = match.group(1)
            break
    
    # Convert correct answer numbers to letters
    correct_letters = ""
    if correct_answer:
        # Parse "1" or "1,4" or "2,3"
        numbers = [int(n.strip()) for n in correct_answer.split(',')]
        correct_letters = ','.join([chr(64 + n) for n in numbers])  # 1->A, 2->B, etc
    
    return {
        'num': q['num'],
        'question': question_text,
        'options': options,
        'explanation': explanation.strip(),
        'correct_answer': correct_letters
    }

# Parse all questions
parsed_questions = []
for q in questions:
    parsed = parse_question_final(q)
    if parsed:
        parsed_questions.append(parsed)

print(f"Successfully parsed {len(parsed_questions)} questions\n")

# Show first 3 for verification
print("Sample parsed questions:")
for i in range(min(3, len(parsed_questions))):
    p = parsed_questions[i]
    print(f"\nQ{p['num']}:")
    print(f"  Question: {p['question'][:80]}...")
    print(f"  Options: {list(p['options'].keys())}")
    print(f"  Correct: {p['correct_answer']}")

# Load the template Excel file
print("\n\nLoading template Excel file...")
excel_path = 'data/DBx_Questions.xlsx'
template_df = pd.read_excel(excel_path)

# Find the max QID in the template
max_qid = template_df['QID'].max()
next_qid = int(max_qid) + 1
print(f"Max QID in template: {max_qid}")
print(f"Starting new QID: {next_qid}")

# Create new rows for the parsed questions
new_rows = []
for i, pq in enumerate(parsed_questions):
    qid = next_qid + i
    
    # Build row dict matching the template columns
    row = {
        'Source': 'SG',
        'QID': qid,
        'Exam': 'DCDEA Practice Exam 2',
        'Number': pq['num'],
        'Topic': '',  # To be filled manually if needed
        'Subtopic': '',
        'Difficulty': '',
        'Question': pq['question'],
        'A': pq['options'].get('A', ''),
        'B': pq['options'].get('B', ''),
        'C': pq['options'].get('C', ''),
        'D': pq['options'].get('D', ''),
        'E': pq['options'].get('E', ''),
        'F': pq['options'].get('F', ''),
        'CorrectLetter': pq['correct_answer'],
        'Explanation': pq['explanation'],
        'Notes': '',
        'LastReviewed': None,
        'CorrectCount': 0,
        'IncorrectCount': 0,
        'FlagForReview': False,
        'Tags': '',
        'Graphic': '',
        'Deeper Dive': '',
        'VerifiedAnswer': '',
        'VerificationStatus': '',
        'AnswerMatchesWorkbook': '',
        'VerificationNotes': '',
        'VerificationSourceKeys': '',
        'VerifiedOn': None,
        'ToDelete': '',
        'DateAdded': datetime.now()
    }
    new_rows.append(row)

# Create DataFrame from new rows
new_df = pd.DataFrame(new_rows)

# Append to template
combined_df = pd.concat([template_df, new_df], ignore_index=True)

# Save to new Excel file
output_path = 'data/DCDEA_Practice_Exam_2_Questions.xlsx'
print(f"\nSaving {len(new_rows)} new questions to {output_path}...")
combined_df.to_excel(output_path, index=False)

print(f"✓ Successfully created {output_path}")
print(f"  Total questions now: {len(combined_df)}")
print(f"  New QIDs: {next_qid} to {next_qid + len(new_rows) - 1}")

