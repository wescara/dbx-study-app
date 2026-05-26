from docx import Document
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

docx_path = 'data/DCDEA Practice Exam 2.docx'
doc = Document(docx_path)

# Parse questions from paragraphs
questions = []
current_question = None
current_lines = []

for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        continue
    
    # Check for question header
    question_match = re.match(r'^Question\s+(\d+):', text)
    if question_match:
        if current_question is not None:
            current_question['lines'] = current_lines
            questions.append(current_question)
        
        current_question = {'num': int(question_match.group(1))}
        current_lines = [text]
        continue
    
    if current_question is not None:
        if text.startswith('Correct Anwser:') or text.startswith('Correct Answer:'):
            current_lines.append(text)
            current_question['lines'] = current_lines
            questions.append(current_question)
            current_question = None
            current_lines = []
            continue
        
        current_lines.append(text)

if current_question is not None:
    current_question['lines'] = current_lines
    questions.append(current_question)

print(f"Total questions: {len(questions)}\n")

# Parse each question into components
def parse_question(q):
    lines = q['lines']
    
    # Find the main question text (usually line 1)
    question_text = lines[1] if len(lines) > 1 else ""
    
    # Find options and explanation
    options = {}
    explanation = ""
    correct_answer = ""
    explanation_start = -1
    option_letters = ['A', 'B', 'C', 'D', 'E', 'F']
    current_option = 0
    
    for i, line in enumerate(lines[2:], start=2):
        # Look for separator line
        if line.startswith('-'):
            explanation_start = i + 1
            break
        
        # Try to detect if this is an option
        # Could be "(N) text" or just text
        if current_option < len(option_letters):
            # Check if it looks like a numbered option
            if re.match(r'^\(\d+\)', line) or current_option == 0:
                # This is an option
                option_text = re.sub(r'^\(\d+\)\s*', '', line).strip()
                options[option_letters[current_option]] = option_text
                current_option += 1
            else:
                # Might be continuation or something else
                options[option_letters[current_option]] = line
                current_option += 1
    
    # Extract explanation
    if explanation_start >= 0 and explanation_start < len(lines):
        explanation_lines = []
        for line in lines[explanation_start:]:
            if 'Correct Anwser:' in line or 'Correct Answer:' in line:
                # Extract correct answer
                match = re.search(r'\(([^)]+)\)', line)
                if match:
                    correct_answer = match.group(1)
                break
            if not line.startswith('-'):
                if 'Documentation' not in line and 'Reference' not in line and 'Managed connectors' not in line and 'pyspark' not in line and 'Pipelines Mode' not in line and 'What is' not in line:
                    explanation_lines.append(line)
        
        explanation = ' '.join(explanation_lines).strip()
    
    return {
        'qid_num': q['num'],
        'question': question_text,
        'options': options,
        'explanation': explanation,
        'correct_answer': correct_answer
    }

# Parse all questions
parsed_questions = []
for q in questions:
    parsed_q = parse_question(q)
    parsed_questions.append(parsed_q)
    
    if parsed_q['qid_num'] <= 3:
        print(f"Q{parsed_q['qid_num']}:")
        print(f"  Question: {parsed_q['question'][:80]}...")
        print(f"  Options: {parsed_q['options']}")
        print(f"  Correct: {parsed_q['correct_answer']}")
        print(f"  Explanation: {parsed_q['explanation'][:80]}...")
        print()

print(f"\nParsed {len(parsed_questions)} questions successfully")

# Save parsed data for verification
with open('parsed_exam_data.txt', 'w', encoding='utf-8') as f:
    for pq in parsed_questions:
        f.write(f"\n{'='*80}\n")
        f.write(f"Question {pq['qid_num']}\n")
        f.write(f"{'='*80}\n")
        f.write(f"Question: {pq['question']}\n\n")
        f.write("Options:\n")
        for letter, text in pq['options'].items():
            f.write(f"  {letter}: {text}\n")
        f.write(f"\nCorrect Answer: {pq['correct_answer']}\n")
        f.write(f"\nExplanation: {pq['explanation']}\n")

print("Saved parsed data to parsed_exam_data.txt")
